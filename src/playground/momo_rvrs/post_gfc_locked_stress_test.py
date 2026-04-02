from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from src.playground.momo_rvrs.data_loader import (
    get_benchmark_ohlc,
    load_ftse250_open_close_from_parquet,
    load_ftse250_precomputed_returns,
)
from src.playground.momo_rvrs.locked_strategy import (
    apply_total_cost_on_rebalance_days,
    corr_and_beta,
    drawdown_curve,
    enforce_dollar_neutral_on_rebalance,
    equity_curve_from_log_returns,
    gross_exposures,
    make_comparison_plots,
    portfolio_log_returns_same_day,
    print_period_compact_table,
    rolling_sharpe,
    tilted_buffered_long_short_weights,
    turnover_on_rebalance,
)
from src.playground.momo_rvrs.signals import (
    forward_fill_weights,
    get_monthly_rebalance_dates,
    perf_metrics_from_log_returns,
    vol_adjusted_formation_signal,
)

# ============================================================
# Locked strategy config
# ============================================================

UNIVERSE = "FTSE250"
START = "2009-07-01"
END = None

FORMATION_WINDOW = 10
VOL_WINDOW = 20
REBALANCE = "monthly"
MIN_NAMES = 200

LONG_ENTRY = 0.10
LONG_EXIT = 0.35
SHORT_ENTRY = 0.90
SHORT_EXIT = 0.75

TARGET_GROSS = 2.0
USE_EXPOSURE_SCALING = False
LAMBDA_EARN_MONTHS = 1.25
LAMBDA_OTHER_MONTHS = 1.0
EARN_MONTHS = {1, 2, 4, 5, 7, 8, 10, 11}

BASE_TC_BPS = 25.0
BASE_SLIPPAGE_BPS = 10.0

OUT_DIR = Path("src/playground/momo_rvrs/post_gfc_locked_outputs")
OUT_DIR.mkdir(parents=True, exist_ok=True)




# ============================================================
# Data classes
# ============================================================

@dataclass
class RunResult:
    name: str
    gross_log: pd.Series
    net_log: pd.Series
    bench_log: pd.Series
    weights_on_reb: pd.DataFrame
    reb_turn: pd.Series
    gross_metrics: dict
    net_metrics: dict
    corr_to_bench: float
    beta_to_bench: float
    avg_reb_turnover: float
    avg_gross_long: float
    avg_gross_short: float
    avg_net: float
    final_equity_net: float


# ============================================================
# Helpers
# ============================================================

def prepare_post_gfc_data() -> tuple[pd.DataFrame, pd.DataFrame, pd.Series]:
    open_px, close_px = load_ftse250_open_close_from_parquet(start=START, end=END)
    cc_log, gap_log = load_ftse250_precomputed_returns(start=START, end=END)

    bench = get_benchmark_ohlc(UNIVERSE, start=START, end=END, force=False)
    bench_log = np.log(bench["Close"]).diff()

    idx = gap_log.index.intersection(cc_log.index).intersection(bench_log.index)
    gap_log = gap_log.loc[idx]
    cc_log = cc_log.loc[idx]
    bench_log = bench_log.loc[idx]

    return gap_log, cc_log, bench_log


def build_locked_weights(gap_log: pd.DataFrame, cc_log: pd.DataFrame) -> pd.DataFrame:
    signal = vol_adjusted_formation_signal(
        log_rets=cc_log,
        formation_window=FORMATION_WINDOW,
        vol_window=VOL_WINDOW,
        vol_floor=1e-6,
    )

    valid_dates = gap_log.index[gap_log.notna().any(axis=1)]
    signal = signal.loc[valid_dates]

    tradable_counts = signal.notna().sum(axis=1)
    ok_dates = tradable_counts[tradable_counts >= MIN_NAMES].index
    signal = signal.loc[ok_dates]

    rebalance_dates = get_monthly_rebalance_dates(signal.index)

    w_on_reb = tilted_buffered_long_short_weights(
        signal=signal,
        rebalance_dates=rebalance_dates,
        long_entry=LONG_ENTRY,
        long_exit=LONG_EXIT,
        short_entry=SHORT_ENTRY,
        short_exit=SHORT_EXIT,
        min_names=MIN_NAMES,
        ew_within_side=True,
    )

    w_on_reb = enforce_dollar_neutral_on_rebalance(w_on_reb, target_gross=TARGET_GROSS)

    if USE_EXPOSURE_SCALING:
        w_on_reb = apply_exposure_scaling_on_rebalance(
            weights_on_reb=w_on_reb,
            earn_months=EARN_MONTHS,
            lambda_earn_months=LAMBDA_EARN_MONTHS,
            lambda_other_months=LAMBDA_OTHER_MONTHS,
        )

    return w_on_reb


def evaluate_with_costs(
    name: str,
    gap_log: pd.DataFrame,
    bench_log: pd.Series,
    weights_on_reb: pd.DataFrame,
    tc_bps: float,
    slippage_bps: float,
) -> RunResult:
    w_daily = forward_fill_weights(weights_on_reb).reindex(gap_log.index).ffill().fillna(0.0)

    gross = portfolio_log_returns_same_day(gap_log, w_daily)
    reb_turn = turnover_on_rebalance(weights_on_reb)

    net = apply_total_cost_on_rebalance_days(
        gross,
        reb_turn,
        tc_bps_per_1turn=tc_bps,
        close_slippage_bps_per_1turn=slippage_bps,
    )

    common = net.index.intersection(bench_log.index)
    gross = gross.loc[common].dropna()
    net = net.loc[common].dropna()
    bench = bench_log.loc[common].dropna()

    common2 = gross.index.intersection(net.index).intersection(bench.index)
    gross = gross.loc[common2]
    net = net.loc[common2]
    bench = bench.loc[common2]

    gross_m = perf_metrics_from_log_returns(gross)
    net_m = perf_metrics_from_log_returns(net)
    corr, beta = corr_and_beta(net, bench)

    eq_net = equity_curve_from_log_returns(net)
    gl, gs, n = gross_exposures(w_daily.reindex(common2))

    return RunResult(
        name=name,
        gross_log=gross,
        net_log=net,
        bench_log=bench,
        weights_on_reb=weights_on_reb,
        reb_turn=reb_turn,
        gross_metrics=gross_m,
        net_metrics=net_m,
        corr_to_bench=corr,
        beta_to_bench=beta,
        avg_reb_turnover=float(reb_turn.mean()) if len(reb_turn) else np.nan,
        avg_gross_long=float(gl.mean()) if len(gl) else np.nan,
        avg_gross_short=float(gs.mean()) if len(gs) else np.nan,
        avg_net=float(n.mean()) if len(n) else np.nan,
        final_equity_net=float(eq_net.iloc[-1]) if len(eq_net) else np.nan,
    )


def summarize_by_period(strategy_name: str, strategy_log: pd.Series, bench_log: pd.Series) -> pd.DataFrame:
    periods = [
        ("Full_Post_GFC", pd.Timestamp("2009-07-01"), pd.Timestamp("2099-12-31")),
        ("Recovery_2009_2012", pd.Timestamp("2009-07-01"), pd.Timestamp("2012-12-31")),
        ("QE_2013_2019", pd.Timestamp("2013-01-01"), pd.Timestamp("2019-12-31")),
        ("COVID_2020", pd.Timestamp("2020-01-01"), pd.Timestamp("2020-12-31")),
        ("Inflation_2021_now", pd.Timestamp("2021-01-01"), pd.Timestamp("2099-12-31")),
    ]

    rows = []
    for period_name, start, end in periods:
        s = strategy_log.loc[start:end].dropna()
        b = bench_log.loc[start:end].dropna()

        common = s.index.intersection(b.index)
        s = s.loc[common]
        b = b.loc[common]

        if len(common) < 20:
            continue

        m = perf_metrics_from_log_returns(s)
        corr, beta = corr_and_beta(s, b)
        eq = equity_curve_from_log_returns(s)

        rows.append(
            {
                "strategy": strategy_name,
                "period": period_name,
                "start": common.min().date().isoformat(),
                "end": common.max().date().isoformat(),
                "n_days": int(len(common)),
                "sharpe": float(m.get("sharpe", np.nan)),
                "ann_return": float(m.get("ann_return", np.nan)),
                "ann_vol": float(m.get("ann_vol", np.nan)),
                "max_drawdown": float(m.get("max_drawdown", np.nan)),
                "beta_to_benchmark": float(beta),
                "corr_to_benchmark": float(corr),
                "final_equity": float(eq.iloc[-1]),
            }
        )

    return pd.DataFrame(rows)


def print_run_summary(run: RunResult) -> None:
    print(f"\n--- {run.name} ---")
    print(f"Avg TURNOVER per REBALANCE: {run.avg_reb_turnover:.4f}")
    print(f"Avg gross long:            {run.avg_gross_long:.3f}")
    print(f"Avg gross short:           {run.avg_gross_short:.3f}")
    print(f"Avg net:                   {run.avg_net:.3f}")
    print("GROSS:")
    print(f"  Sharpe:    {run.gross_metrics['sharpe']:.4f}")
    print(f"  AnnRet:    {run.gross_metrics['ann_return']:.4f}")
    print(f"  AnnVol:    {run.gross_metrics['ann_vol']:.4f}")
    print(f"  MaxDD:     {run.gross_metrics['max_drawdown']:.4f}")
    print("NET:")
    print(f"  Sharpe:    {run.net_metrics['sharpe']:.4f}")
    print(f"  AnnRet:    {run.net_metrics['ann_return']:.4f}")
    print(f"  AnnVol:    {run.net_metrics['ann_vol']:.4f}")
    print(f"  MaxDD:     {run.net_metrics['max_drawdown']:.4f}")
    print(f"  Final Eq:  {run.final_equity_net:.4f}")
    print("vs Benchmark:")
    print(f"  Corr:      {run.corr_to_bench:.4f}")
    print(f"  Beta:      {run.beta_to_bench:.4f}")


def print_cost_summary(df: pd.DataFrame) -> None:
    print("\n" + "=" * 90)
    print("COST STRESS TEST | POST-GFC LOCKED STRATEGY")
    print("=" * 90)
    for _, row in df.iterrows():
        print(
            f"{row['scenario']:>15} | "
            f"tc={row['tc_bps']:>5.1f} | "
            f"slip={row['slippage_bps']:>5.1f} | "
            f"Sharpe={row['net_sharpe']:>6.3f} | "
            f"AnnRet={row['net_annret']:>6.3f} | "
            f"AnnVol={row['net_annvol']:>6.3f} | "
            f"MaxDD={row['net_maxdd']:>7.3f} | "
            f"Beta={row['beta_to_bench']:>7.3f}"
        )


def print_period_summary(df: pd.DataFrame) -> None:
    print("\n" + "=" * 90)
    print("SUBPERIOD STABILITY | POST-GFC LOCKED STRATEGY")
    print("=" * 90)
    for _, row in df.iterrows():
        print(
            f"{row['period']:>18} | "
            f"{row['start']} -> {row['end']} | "
            f"n={int(row['n_days']):>4d} | "
            f"Sharpe={row['sharpe']:>6.3f} | "
            f"AnnRet={row['ann_return']:>6.3f} | "
            f"AnnVol={row['ann_vol']:>6.3f} | "
            f"MaxDD={row['max_drawdown']:>7.3f} | "
            f"Beta={row['beta_to_benchmark']:>7.3f}"
        )


def save_turnover_diagnostics(run: RunResult) -> pd.DataFrame:
    reb_turn = run.reb_turn.dropna()
    out = pd.DataFrame(
        {
            "rebalance_date": reb_turn.index,
            "turnover": reb_turn.values,
        }
    )
    out.to_csv(OUT_DIR / "locked_post_gfc_rebalance_turnover.csv", index=False)

    summary = pd.DataFrame(
        [
            {
                "avg_reb_turnover": float(reb_turn.mean()),
                "median_reb_turnover": float(reb_turn.median()),
                "p90_reb_turnover": float(reb_turn.quantile(0.90)),
                "max_reb_turnover": float(reb_turn.max()),
                "n_rebalances": int(len(reb_turn)),
            }
        ]
    )
    summary.to_csv(OUT_DIR / "locked_post_gfc_turnover_summary.csv", index=False)
    return summary

PMC_COLORS = {
    "primary": "#1F3044",
    "secondary": "#555555",
    "accent": "#97C5EB",
    "dark_neutral": "#222A35",
    "light_neutral": "#F0F3F5",
    "accent_2": "#3B6189",
}


def apply_pmc_plot_style() -> None:
    """
    Apply PMC plot styling to matplotlib.
    """
    available_fonts = {f.name for f in font_manager.fontManager.ttflist}
    chosen_font = "Source Sans Pro" if "Source Sans Pro" in available_fonts else "DejaVu Sans"

    plt.rcParams.update(
        {
            "font.family": chosen_font,
            "axes.titlesize": 14,
            "axes.labelsize": 11,
            "xtick.labelsize": 10,
            "ytick.labelsize": 10,
            "legend.fontsize": 10,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.grid": True,
            "grid.alpha": 0.25,
            "grid.linestyle": "--",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
            "savefig.bbox": "tight",
        }
    )


def apply_exposure_scaling_on_rebalance(
    weights_on_reb: pd.DataFrame,
    earn_months: set[int],
    lambda_earn_months: float,
    lambda_other_months: float,
) -> pd.DataFrame:
    """
    Scale rebalance weights by month-level leverage regime.

    Earnings-heavy months get higher exposure.
    Other months get lower exposure.
    """
    w = weights_on_reb.copy()

    for dt in w.index:
        lam = lambda_earn_months if dt.month in earn_months else lambda_other_months
        w.loc[dt] = w.loc[dt] * lam

    return w


def save_results_to_excel(
    out_path: Path,
    base_run: RunResult,
    cost_df: pd.DataFrame,
    period_df: pd.DataFrame,
    turnover_summary: pd.DataFrame,
) -> None:
    """
    Save strategy outputs to a formatted Excel workbook with multiple sheets.
    """
    base_summary = pd.DataFrame(
        [
            {
                "strategy": base_run.name,
                "gross_sharpe": float(base_run.gross_metrics.get("sharpe", np.nan)),
                "gross_annret": float(base_run.gross_metrics.get("ann_return", np.nan)),
                "gross_annvol": float(base_run.gross_metrics.get("ann_vol", np.nan)),
                "gross_maxdd": float(base_run.gross_metrics.get("max_drawdown", np.nan)),
                "net_sharpe": float(base_run.net_metrics.get("sharpe", np.nan)),
                "net_annret": float(base_run.net_metrics.get("ann_return", np.nan)),
                "net_annvol": float(base_run.net_metrics.get("ann_vol", np.nan)),
                "net_maxdd": float(base_run.net_metrics.get("max_drawdown", np.nan)),
                "beta_to_bench": float(base_run.beta_to_bench),
                "corr_to_bench": float(base_run.corr_to_bench),
                "avg_reb_turnover": float(base_run.avg_reb_turnover),
                "avg_gross_long": float(base_run.avg_gross_long),
                "avg_gross_short": float(base_run.avg_gross_short),
                "avg_net": float(base_run.avg_net),
                "final_equity_net": float(base_run.final_equity_net),
            }
        ]
    )

    timeseries_df = pd.DataFrame(
        {
            "strategy_log": base_run.net_log,
            "benchmark_log": base_run.bench_log,
            "strategy_eq": equity_curve_from_log_returns(base_run.net_log),
            "benchmark_eq": equity_curve_from_log_returns(base_run.bench_log),
            "strategy_dd": drawdown_curve(equity_curve_from_log_returns(base_run.net_log)),
            "benchmark_dd": drawdown_curve(equity_curve_from_log_returns(base_run.bench_log)),
        }
    )

    turnover_detail = pd.DataFrame(
        {
            "rebalance_date": base_run.reb_turn.index,
            "turnover": base_run.reb_turn.values,
        }
    )

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        base_summary.to_excel(writer, sheet_name="summary", index=False)
        cost_df.to_excel(writer, sheet_name="cost_stress", index=False)
        period_df.to_excel(writer, sheet_name="subperiods", index=False)
        turnover_summary.to_excel(writer, sheet_name="turnover_summary", index=False)
        turnover_detail.to_excel(writer, sheet_name="turnover_detail", index=False)
        timeseries_df.to_excel(writer, sheet_name="timeseries")

        wb = writer.book
        header_fill = PatternFill(fill_type="solid", fgColor="1F3044")
        header_font = Font(color="FFFFFF", bold=True)
        center_align = Alignment(horizontal="center", vertical="center")

        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = center_align

            for col_cells in ws.columns:
                max_len = 0
                col_idx = col_cells[0].column
                col_letter = get_column_letter(col_idx)

                for cell in col_cells:
                    val = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, len(val))

                ws.column_dimensions[col_letter].width = min(max_len + 2, 24)


# ============================================================
# Main
# ============================================================

def main() -> None:
    print("\n" + "=" * 90)
    print("LOCKED STRATEGY STRESS TEST | POST-GFC")
    print("=" * 90)
    print(f"Universe:           {UNIVERSE}")
    print(f"Sample:             {START} -> latest")
    print(f"Formation window:   {FORMATION_WINDOW}")
    print(f"Thresholds:         LE={LONG_ENTRY:.2f} LX={LONG_EXIT:.2f} SE={SHORT_ENTRY:.2f} SX={SHORT_EXIT:.2f}")
    print(f"Exposure scaling:   {USE_EXPOSURE_SCALING}")
    if USE_EXPOSURE_SCALING:
        print(f"Lambda earn months: {LAMBDA_EARN_MONTHS}")
        print(f"Lambda other months:{LAMBDA_OTHER_MONTHS}")
        print(f"Earnings months:    {sorted(EARN_MONTHS)}")
    print(f"Min names:          {MIN_NAMES}")
    print(f"Target gross:       {TARGET_GROSS}")

    gap_log, cc_log, bench_log = prepare_post_gfc_data()
    weights_on_reb = build_locked_weights(gap_log, cc_log)

    base_run = evaluate_with_costs(
        name="locked_post_gfc_base",
        gap_log=gap_log,
        bench_log=bench_log,
        weights_on_reb=weights_on_reb,
        tc_bps=BASE_TC_BPS,
        slippage_bps=BASE_SLIPPAGE_BPS,
    )
    print_run_summary(base_run)

    turnover_summary = save_turnover_diagnostics(base_run)

    make_comparison_plots(base_run.net_log, base_run.bench_log)

    base_ts = pd.DataFrame(
        {
            "strategy_log": base_run.net_log,
            "benchmark_log": base_run.bench_log,
            "strategy_eq": equity_curve_from_log_returns(base_run.net_log),
            "benchmark_eq": equity_curve_from_log_returns(base_run.bench_log),
        }
    )
    base_ts.to_csv(OUT_DIR / "locked_post_gfc_timeseries.csv")

    scenarios = [
        ("base_realistic", 10.0, 10.0),
        ("tc_25bps", 25.0, 10.0),
        ("tc_50bps", 50.0, 10.0),
        ("tc_75bps", 75.0, 10.0),
        ("tc_100bps", 100.0, 10.0),
    ]

    rows = []
    for scenario_name, tc_bps, slip_bps in scenarios:
        run = evaluate_with_costs(
            name=scenario_name,
            gap_log=gap_log,
            bench_log=bench_log,
            weights_on_reb=weights_on_reb,
            tc_bps=tc_bps,
            slippage_bps=slip_bps,
        )
        rows.append(
            {
                "scenario": scenario_name,
                "tc_bps": tc_bps,
                "slippage_bps": slip_bps,
                "net_sharpe": float(run.net_metrics.get("sharpe", np.nan)),
                "net_annret": float(run.net_metrics.get("ann_return", np.nan)),
                "net_annvol": float(run.net_metrics.get("ann_vol", np.nan)),
                "net_maxdd": float(run.net_metrics.get("max_drawdown", np.nan)),
                "beta_to_bench": float(run.beta_to_bench),
                "corr_to_bench": float(run.corr_to_bench),
                "avg_reb_turnover": float(run.avg_reb_turnover),
                "final_equity_net": float(run.final_equity_net),
            }
        )

    cost_df = pd.DataFrame(rows)
    cost_df.to_csv(OUT_DIR / "locked_post_gfc_cost_stress.csv", index=False)
    print_cost_summary(cost_df)

    period_df = summarize_by_period(
        strategy_name="locked_post_gfc_base",
        strategy_log=base_run.net_log,
        bench_log=base_run.bench_log,
    )
    period_df.to_csv(OUT_DIR / "locked_post_gfc_subperiods.csv", index=False)
    print_period_summary(period_df)

    excel_path = OUT_DIR / "locked_post_gfc_results.xlsx"
    save_results_to_excel(
        out_path=excel_path,
        base_run=base_run,
        cost_df=cost_df,
        period_df=period_df,
        turnover_summary=turnover_summary,
    )

    print("\n" + "=" * 90)
    print("SAVED OUTPUTS")
    print("=" * 90)
    print(f"  {OUT_DIR / 'locked_post_gfc_timeseries.csv'}")
    print(f"  {OUT_DIR / 'locked_post_gfc_cost_stress.csv'}")
    print(f"  {OUT_DIR / 'locked_post_gfc_subperiods.csv'}")
    print(f"  {OUT_DIR / 'locked_post_gfc_rebalance_turnover.csv'}")
    print(f"  {OUT_DIR / 'locked_post_gfc_turnover_summary.csv'}")
    print(f"  {OUT_DIR / 'equity_curve_strategy_vs_ftse250.png'}")
    print(f"  {OUT_DIR / 'drawdown_strategy_vs_ftse250.png'}")
    print(f"  {OUT_DIR / 'rolling_vol_strategy_vs_ftse250.png'}")
    print(f"  {OUT_DIR / 'rolling_sharpe_strategy_vs_ftse250.png'}")
    print(f"  {OUT_DIR / 'locked_post_gfc_results.xlsx'}")
    print("=" * 90)


if __name__ == "__main__":
    main()

